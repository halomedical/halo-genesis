"""OCR and document text extraction for special upload workflows."""

from __future__ import annotations

import logging

from io import BytesIO
from pathlib import Path

from halo_sentry.privacy import safe_file_ref

from halo_sentry.config import SpecialWorkflowConfig

logger = logging.getLogger("halo_sentry.text_extraction")


class TextExtractionError(RuntimeError):
    """Base class for text extraction failures."""


class TextExtractionUnavailable(TextExtractionError):
    """Raised when an optional OCR/document dependency is unavailable."""


class UnsupportedFileType(TextExtractionError):
    """Raised when a file cannot be interpreted by the selected extractor."""


class TextExtractor:
    def __init__(self, config: SpecialWorkflowConfig) -> None:
        self._config = config

    def validate_dependencies(self) -> None:
        try:
            from PIL import Image  # noqa: F401
            import pytesseract
        except ImportError as exc:
            raise TextExtractionUnavailable(
                "Image identifier checking requires Pillow and pytesseract."
            ) from exc

        self._configure_tesseract(pytesseract)
        try:
            pytesseract.get_tesseract_version()
        except Exception as exc:
            raise TextExtractionUnavailable(
                "Tesseract OCR is not available. Install Tesseract for "
                "Windows or configure special_workflows.ocr_tesseract_cmd."
            ) from exc

    def extract_image_text(self, path: Path) -> str:
        try:
            from PIL import Image, UnidentifiedImageError
            import pytesseract
            from pytesseract.pytesseract import (
                TesseractError,
                TesseractNotFoundError,
            )
        except ImportError as exc:
            raise TextExtractionUnavailable(
                "Image OCR requires Pillow and pytesseract."
            ) from exc

        self._configure_tesseract(pytesseract)

        try:
            with Image.open(path) as image:
                image.load()
                return pytesseract.image_to_string(image)
        except (TesseractError, TesseractNotFoundError) as exc:
            raise TextExtractionUnavailable(
                "Tesseract OCR failed while reading image."
            ) from exc
        except (UnidentifiedImageError, OSError) as exc:
            raise UnsupportedFileType(
                "File is not a supported image"
            ) from exc

    def extract_pdf_text(self, path: Path) -> str:
        try:
            from PIL import Image
            import fitz
            import pytesseract
            from pytesseract.pytesseract import (
                TesseractError,
                TesseractNotFoundError,
            )
        except ImportError as exc:
            raise TextExtractionUnavailable(
                "PDF OCR requires Pillow, PyMuPDF, and pytesseract."
            ) from exc

        self._configure_tesseract(pytesseract)
        text_parts: list[str] = []
        try:
            with fitz.open(str(path)) as document:
                page_count = min(
                    len(document),
                    self._config.pdf_ocr_max_pages,
                )
                for page_index in range(page_count):
                    page = document.load_page(page_index)
                    embedded_text = page.get_text("text")
                    if embedded_text.strip():
                        text_parts.append(embedded_text)

                    matrix = fitz.Matrix(2, 2)
                    pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                    image = Image.frombytes(
                        "RGB",
                        (pixmap.width, pixmap.height),
                        pixmap.samples,
                    )
                    try:
                        text_parts.append(pytesseract.image_to_string(image))
                    except (TesseractError, TesseractNotFoundError) as exc:
                        raise TextExtractionUnavailable(
                            "Tesseract OCR failed while reading PDF."
                        ) from exc
        except Exception as exc:
            if isinstance(exc, TextExtractionUnavailable):
                raise
            raise UnsupportedFileType(
                "File is not a supported PDF"
            ) from exc

        return "\n".join(text_parts)

    def extract_docx_text(self, path: Path) -> str:
        try:
            from PIL import Image, UnidentifiedImageError
            import docx
            import pytesseract
            from pytesseract.pytesseract import (
                TesseractError,
                TesseractNotFoundError,
            )
        except ImportError as exc:
            raise TextExtractionUnavailable(
                "DOCX extraction requires python-docx, Pillow, and pytesseract."
            ) from exc

        self._configure_tesseract(pytesseract)
        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise UnsupportedFileType(
                "File is not a supported DOCX document"
            ) from exc

        text_parts: list[str] = []
        text_parts.extend(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        for part in document.part.related_parts.values():
            if not part.content_type.startswith("image/"):
                continue
            try:
                with Image.open(BytesIO(part.blob)) as image:
                    image.load()
                    text_parts.append(pytesseract.image_to_string(image))
            except (TesseractError, TesseractNotFoundError) as exc:
                raise TextExtractionUnavailable(
                    "Tesseract OCR failed while reading DOCX image."
                ) from exc
            except (UnidentifiedImageError, OSError):
                logger.debug("Skipping unsupported image inside DOCX: %s", safe_file_ref(path))

        return "\n".join(text_parts)

    def _configure_tesseract(self, pytesseract_module: object) -> None:
        if self._config.ocr_tesseract_cmd:
            pytesseract_module.pytesseract.tesseract_cmd = (
                self._config.ocr_tesseract_cmd
            )
